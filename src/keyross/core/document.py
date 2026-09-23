"""The canonical document: what oracles reason about. Rows with stable identifiers (rid:)."""
from __future__ import annotations

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable

DEFAULT_MAPPING = {
    "designation": ["designation", "désignation", "libelle", "libellé", "description"],
    "unit": ["unite", "unité", "u", "unit"],
    "qty": ["quantite", "quantité", "qte", "qté", "qty"],
    "unit_price": ["pu", "prix unitaire", "p.u.", "unit_price", "prix_unitaire"],
    "amount": ["montant", "total", "amount", "montant ht"],
    "number": ["n°", "no", "num", "numero", "numéro", "ref", "code"],
}


@dataclass
class Line:
    rid: str                       # stable identifier: rid:<source row number>
    row: int                       # source row (1-based)
    number: str | None = None      # hierarchical numbering (1.2.3)
    designation: str = ""
    unit: str | None = None
    qty: float | None = None
    unit_price: float | None = None
    amount: float | None = None
    raw: dict[str, Any] = field(default_factory=dict)

    @property
    def is_subtotal(self) -> bool:
        d = self.designation.strip().lower()
        return d.startswith(("sous-total", "sous total", "total")) and self.amount is not None and self.qty is None

    @property
    def is_amount(self) -> bool:
        """A priced line: it carries an amount or a quantity / unit price pair, and is not a subtotal."""
        return (self.amount is not None or (self.qty is not None and self.unit_price is not None)) and not self.is_subtotal


@dataclass
class Document:
    path: str
    lines: list[Line]
    columns: dict[str, str] = field(default_factory=dict)    # canonical field -> source header
    meta: dict[str, Any] = field(default_factory=dict)       # doc_type, contractor, lot… (provided by the context)

    def amount_lines(self) -> list[Line]:
        return [l for l in self.lines if l.is_amount]

    def row(self, rid: str) -> Line | None:
        return next((l for l in self.lines if l.rid == rid), None)

    def blocks(self) -> list[tuple[list[Line], Line]]:
        """Split into blocks (priced lines, followed by their subtotal line)."""
        out, current = [], []
        for l in self.lines:
            if l.is_subtotal:
                out.append((current, l)); current = []
            elif l.is_amount:
                current.append(l)
        return out


def _to_float(v: Any) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\u202f", "").replace(" ", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def _map_headers(headers: Iterable[Any], mapping: dict[str, list[str]]) -> dict[str, int]:
    idx: dict[str, int] = {}
    for i, h in enumerate(headers):
        if h is None:
            continue
        key = str(h).strip().lower()
        for field_name, aliases in mapping.items():
            if field_name not in idx and key in aliases:
                idx[field_name] = i
    return idx


def _rows_to_document(path: str, rows: list[list[Any]], mapping: dict[str, list[str]] | None = None) -> Document:
    mapping = mapping or DEFAULT_MAPPING
    header_row, idx = None, {}
    for r, row in enumerate(rows[:30]):          # the header: first row with at least two known fields
        m = _map_headers(row, mapping)
        if len(m) >= 2:
            header_row, idx = r, m
            break
    if header_row is None:
        raise ValueError(f"{path}: no header row recognized (expected at least two columns among {list(mapping)})")
    lines: list[Line] = []
    for r in range(header_row + 1, len(rows)):
        row = rows[r]
        if not any(c not in (None, "") for c in row):
            continue

        def get(k: str) -> Any:
            return row[idx[k]] if k in idx and idx[k] < len(row) else None

        lines.append(Line(
            rid=f"rid:{r + 1}", row=r + 1,
            number=str(get("number")).strip() if get("number") not in (None, "") else None,
            designation=str(get("designation") or "").strip(),
            unit=str(get("unit")).strip() if get("unit") not in (None, "") else None,
            qty=_to_float(get("qty")), unit_price=_to_float(get("unit_price")), amount=_to_float(get("amount")),
            raw={str(rows[header_row][i]): row[i] for i in range(min(len(row), len(rows[header_row]))) if rows[header_row][i] is not None},
        ))
    columns = {k: str(rows[header_row][i]) for k, i in idx.items()}
    return Document(path=path, lines=lines, columns=columns)


def load(path: str | Path, sheet: str | None = None, mapping: dict[str, list[str]] | None = None) -> Document:
    """Load an .xlsx or .csv file into the canonical document."""
    p = Path(path)
    if p.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
        try:
            ws = wb[sheet] if sheet else wb.worksheets[0]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
        finally:
            wb.close()                                   # read-only workbooks keep the file open until closed
        return _rows_to_document(str(p), rows, mapping)
    if p.suffix.lower() == ".csv":
        with open(p, newline="", encoding="utf-8-sig") as f:
            rows = [r for r in csv.reader(f, delimiter=";")]
        return _rows_to_document(str(p), rows, mapping)
    raise ValueError(f"unsupported format: {p.suffix}")
